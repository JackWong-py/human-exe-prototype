from dataclasses import dataclass
from pathlib import Path

from docx import Document


@dataclass
class RawDoc:
    doc_type: str
    title: str
    pairs: list
    error: str | None = None
    noisy: bool = False


def detect_doc_type(title: str) -> str:
    """
    Detect document type from the document title.
    """
    title_upper = str(title).upper()

    if "INSTRUCTION" in title_upper:
        return "SI"

    if "BILL OF LADING" in title_upper:
        return "BL"

    if (
        "COMMERCIAL INVOICE" in title_upper
        or "PACKING LIST" in title_upper
        or "CERTIFICATE OF ORIGIN" in title_upper
    ):
        return "OTHER"

    return "OTHER"


def clean_text(text: str) -> str:
    """
    Clean text from a DOCX paragraph or table cell.
    Multiple lines are joined into one value.
    """
    lines = []

    for line in str(text).splitlines():
        line = line.strip()

        if line:
            lines.append(line)

    return " ".join(lines).strip()


def read_docx(path: str | Path) -> RawDoc:
    """
    Read a DOCX document.

    The document may contain:
    - Paragraphs with label/value information
    - Tables containing label/value information
    - Multiple lines inside table cells
    """
    path = Path(path)

    try:
        document = Document(path)
    except Exception:
        return RawDoc(
            doc_type="OTHER",
            title="",
            pairs=[],
            error="unreadable",
            noisy=False,
        )

    # Get the document title from the first non-empty paragraph.
    title = ""

    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)

        if text:
            title = text
            break

    doc_type = detect_doc_type(title)

    pairs = []

    # Read useful label/value information from paragraphs.
    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)

        if not text:
            continue

        # Skip the title.
        if text == title:
            continue

        # Handle paragraphs such as:
        if ":" in text:
            label, value = text.split(":", 1)

            label = label.strip()
            value = value.strip()

            if label and value:
                pairs.append(
                    (
                        label,
                        value,
                        text,
                    )
                )

    # Read label/value information from tables.
    for table in document.tables:
        for row in table.rows:
            cells = row.cells

            if len(cells) < 2:
                continue

            label = clean_text(cells[0].text)
            value = clean_text(cells[1].text)

            if not label or not value:
                continue

            source = f"{label}: {value}"

            pairs.append(
                (
                    label,
                    value,
                    source,
                )
            )

    return RawDoc(
        doc_type=doc_type,
        title=title,
        pairs=pairs,
        error=None,
        noisy=False,
    )


def read_document(path: str | Path) -> RawDoc:
    """
    Read a document based on its file extension.
    """
    path = Path(path)
    extension = path.suffix.lower()

    if extension == ".docx":
        return read_docx(path)

    return RawDoc(
        doc_type="OTHER",
        title="",
        pairs=[],
        error="unreadable",
        noisy=False,
    )